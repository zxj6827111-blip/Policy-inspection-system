"""校验区划导入表与系统内的站点配置、机构知识库是否一致。

用法：
    python scripts/check_region_inventory.py [导入表.xlsx]

导入表四列为：省份、城市、区县、街道。只填到「区县」的行表示市政府工作部门，
填了「街道」的行表示该区县下辖的乡镇街道。系统按代码静态配置登记站点与栏目，
本脚本负责回答“表里的每一项都进系统了吗”。
"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook

from app.config import ORDOS_DEPARTMENTS, SCAN_TARGETS, targets_for_site
from app.knowledge.organizations import ORDOS_CITY_DISTRICT, SEED_ORGANIZATIONS


def read_inventory(path: Path) -> tuple[set[str], set[str], set[str]]:
    sheet = load_workbook(path, read_only=True).active
    districts: set[str] = set()
    departments: set[str] = set()
    townships: set[str] = set()
    for row in sheet.iter_rows(min_row=2, values_only=True):
        cells = [str(value).strip() if value is not None else "" for value in row]
        cells += [""] * (4 - len(cells))
        _province, _city, name, township = cells[:4]
        if not name:
            continue
        if township:
            districts.add(name)
            townships.add(f"{name}·{township}")
        else:
            departments.add(name)
    return districts, departments, townships


def main(argv: list[str]) -> int:
    path = Path(argv[1]) if len(argv) > 1 else Path("data/region_import_template.xlsx")
    if not path.exists():
        print(f"导入表不存在：{path}")
        return 2

    districts, departments, townships = read_inventory(path)
    orgs = {item["standard_name"]: item for item in SEED_ORGANIZATIONS}
    known_aliases = {
        alias
        for item in SEED_ORGANIZATIONS
        for alias in [item["standard_name"], *item.get("aliases", [])]
    }
    known_districts = {item["district"] for item in SEED_ORGANIZATIONS if item["district"]}

    problems: list[str] = []
    notes: list[str] = []

    missing_districts = sorted(districts - known_districts)
    for name in missing_districts:
        problems.append(f"区县「{name}」未登记到机构知识库")

    for name in sorted(districts):
        standard = f"鄂尔多斯市{name}人民政府"
        if standard not in orgs:
            problems.append(f"区县「{name}」缺少人民政府机构条目：{standard}")

    for entry in sorted(townships):
        district, township = entry.split("·", 1)
        suffix = "街道办事处" if township.endswith("街道") else "人民政府"
        standard = f"鄂尔多斯市{district}{township}{suffix}"
        if standard not in orgs:
            problems.append(f"乡镇街道「{township}」未入库：{standard}")

    city_aliases = {
        alias
        for item in SEED_ORGANIZATIONS
        if item["district"] == ORDOS_CITY_DISTRICT
        for alias in [item["standard_name"], *item.get("aliases", [])]
    }
    for name in sorted(departments):
        if name not in city_aliases and f"鄂尔多斯{name}" not in city_aliases:
            problems.append(f"部门「{name}」未入库为鄂尔多斯市级机构")

    configured = {name for _slug, name in ORDOS_DEPARTMENTS}
    # 官网子站名与清单名可能不同（清单「市农牧局」，官网「市农牧业局」），
    # 因此按机构知识库里的别名组比对，而不是按字面名比对。
    city_alias_groups = [
        {item["standard_name"], *item.get("aliases", [])}
        for item in SEED_ORGANIZATIONS
        if item["district"] == ORDOS_CITY_DISTRICT
    ]

    def covered_names(name: str) -> set[str]:
        for group in city_alias_groups:
            if name in group or f"鄂尔多斯{name}" in group:
                return group
        return {name}

    for name in sorted(configured):
        if not (covered_names(name) & departments):
            problems.append(f"配置里的部门栏目「{name}」不在导入表清单中")
    for name in sorted(departments - configured):
        if covered_names(name) & configured:
            continue  # 官网子站用了另一个名字（清单「市农牧局」/官网「市农牧业局」）
        if name == "市人民政府办公室":
            notes.append(
                "「市人民政府办公室」官网无部门文件子栏目，其文件在「政策文件」栏目（发文字号鄂府办发）"
            )
        else:
            notes.append(f"「{name}」官网未单独开设部门文件子栏目，其文件不在本次扫描范围内")

    scanned = {target.source_name for target in targets_for_site("ordos_departments")}
    if scanned != {"部门文件"}:
        problems.append("ordos_departments 站点的来源名不是「部门文件」")
    missing_targets = [
        slug for slug, _name in ORDOS_DEPARTMENTS if f"ordos_dept_{slug}" not in SCAN_TARGETS
    ]
    for slug in missing_targets:
        problems.append(f"部门子站目录 {slug} 没有对应的扫描来源")

    ordos_targets = [target for target in SCAN_TARGETS.values() if target.collector_type == "ordos"]
    print(f"导入表：{path}")
    print(f"区县 {len(districts)} 个 · 部门 {len(departments)} 个 · 乡镇街道 {len(townships)} 项")
    print(f"机构知识库 {len(orgs)} 条 · 可识别机构别名 {len(known_aliases)} 个")
    print(f"鄂尔多斯扫描来源 {len(ordos_targets)} 个，覆盖 {len({t.host for t in ordos_targets})} 个域名")
    if notes:
        print(f"\n官网与清单的 {len(notes)} 处已知差异：")
        for item in notes:
            print(f"  · {item}")
    if problems:
        print(f"\n发现 {len(problems)} 处不一致：")
        for item in problems:
            print(f"  - {item}")
        return 1
    print("\n导入表与系统配置、机构知识库一致。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
